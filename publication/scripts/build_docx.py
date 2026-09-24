"""Convert the manuscript Markdown files into Word documents (python-docx).

Supports the subset used in the manuscript: headings (#, ##, ###), paragraphs,
**bold** and *italic* spans, pipe tables, figure lines ![caption](path),
bullet and numbered lists, and page breaks (a line containing only \\pagebreak).

    python publication/scripts/build_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

PUB = Path(__file__).resolve().parents[1]
DOCS = {'manuscript/manuscript.md': 'manuscript/manuscript.docx',
        'supplementary/supplementary_material.md': 'supplementary/supplementary_material.docx',
        'manuscript/cover_letter.md': 'manuscript/cover_letter.docx'}
INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')


def add_runs(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*'):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def shade(cell):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement('w:shd')
    fill.set(qn('w:val'), 'clear')
    fill.set(qn('w:fill'), 'EDEDED')
    props.append(fill)


def add_table(document, rows):
    header, body = rows[0], rows[2:]
    table = document.add_table(rows=1 + len(body), cols=len(header))
    table.style = 'Table Grid'
    for r, row in enumerate([header] + body):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ''
            paragraph = cell.paragraphs[0]
            add_runs(paragraph, value.strip())
            for run in paragraph.runs:
                run.font.size = Pt(8)
                if r == 0:
                    run.bold = True
            if r == 0:
                shade(cell)
    document.add_paragraph()


def convert(source, target):
    document = Document()
    style = document.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    for level in (1, 2, 3):
        document.styles[f'Heading {level}'].font.name = 'Times New Roman'
    lines = source.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == '\\pagebreak':
            document.add_page_break()
        elif line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            document.add_heading(line.lstrip('#').strip(), level=min(level, 3) if level > 1 else 0)
        elif line.startswith('!['):
            caption, path = re.match(r'!\[(.*)\]\((.*)\)', line).groups()
            document.add_picture(str((source.parent / path).resolve()), width=Inches(6.3))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph = document.add_paragraph()
            add_runs(paragraph, caption)
            for run in paragraph.runs:
                run.font.size = Pt(9)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append([c for c in lines[i].strip().strip('|').split('|')])
                i += 1
            add_table(document, rows)
            continue
        elif re.match(r'^(- |\d+\. )', line):
            numbered = line[0].isdigit()
            paragraph = document.add_paragraph(style='List Number' if numbered else 'List Bullet')
            add_runs(paragraph, re.sub(r'^(- |\d+\. )', '', line))
        else:
            text = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r'^(#|!\[|\||- |\d+\. |\\pagebreak)', lines[i + 1]):
                i += 1
                text.append(lines[i].strip())
            add_runs(document.add_paragraph(), ' '.join(text))
        i += 1
    document.save(target)
    print('wrote', target)


def main():
    for source, target in DOCS.items():
        if (PUB / source).exists():
            convert(PUB / source, PUB / target)


if __name__ == '__main__':
    main()
