"""Create the corrected report from completed results; refuses incomplete experiments.

Use the bundled document runtime and render/inspect the resulting DOCX before delivery.
"""
import json
from pathlib import Path
from statistics import mean
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
RESULTS = ROOT/'results/research_v2'


def main():
    evidence = json.loads((RESULTS/'evaluation_summary.json').read_text())
    paired = json.loads((RESULTS/'paired_seed_comparisons.json').read_text())
    stain = json.loads((RESULTS/'stain_reference_sensitivity.json').read_text())
    models = evidence['models']
    expected = {f'v2_{variant}_seed{seed}' for variant in ('sipakmed_only','combined') for seed in (17,43,101)}
    if set(models) != expected or set(stain['models']) != expected:
        raise ValueError('All six model evaluations and stain analyses must complete before report generation')
    if any(len(stain['models'][name].get(dataset, [])) != 3 for name in expected for dataset in ('dense_evaluation','apc_sparse_test')):
        raise ValueError('Incomplete stain reference analysis')
    doc=Document()
    section=doc.sections[0]
    section.page_width,section.page_height=Inches(8.5),Inches(11)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    for style_name in ('Normal','Title','Heading 1','Heading 2'):
        style=doc.styles[style_name]
        style.font.name='Calibri'
        style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal']
    normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',22),('Heading 1',15),('Heading 2',12)]:
        doc.styles[name].font.size=Pt(size)
    doc.add_paragraph('Corrected cervical cell detection research report','Title')
    doc.add_paragraph('Retrospective training and evaluation after the September 2026 audit')
    text=[]
    def paragraph(value):
        doc.add_paragraph(value)
        text.append(value)
    def heading(value):
        doc.add_heading(value,level=1)
        text.append('\n'+value+'\n')
    def table(headers,rows):
        tab=doc.add_table(rows=1,cols=len(headers))
        tab.style='Table Grid'
        for cell,label in zip(tab.rows[0].cells,headers):
            cell.text=label
            for run in cell.paragraphs[0].runs:run.bold=True
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shading)
        repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell,value in zip(tab.add_row().cells,row):cell.text=str(value)
        text.append(' | '.join(headers))
        text.extend(' | '.join(map(str,row)) for row in rows)
    differences=[paired['paired_by_seed'][str(s)]['dense']['delta']['abnormal_recall'] for s in (17,43,101)]
    direction='higher in all three matched seeds' if all(x>0 for x in differences) else 'not consistently higher across all three matched seeds'
    paragraph(f'The evaluation defects identified in the audit have been corrected and six controlled training runs have been evaluated. Combined-source abnormal-cell recall on the reserved dense evaluation fields was {direction}. This report describes a retrospective research result, not clinical screening validation. The small evaluation sample, model-assisted reference annotations and unavailable patient identities limit the strength of inference.')
    heading('What was corrected')
    paragraph('The five-class checkpoint labels Dyskeratotic and Koilocytotic as classes 0 and 1. Earlier evaluators incorrectly treated classes 2 and 3 as abnormal. Shared name-based mappings now prevent that error. Missed cells are included in support counts, unreviewed predictions cannot be imported as verified annotations, and incompatible classification benchmark deltas have been removed. The earlier claims of scientific completion, established screening thresholds, fixed seed-noise limits and demonstrated morphology invariance are withdrawn.')
    heading('Data and annotation provenance')
    paragraph('The user confirmed reviewing and correctly annotating all 40 dense fields, including checking for cells missed by model proposals. The attestation is tied to the annotation export hash. There are 1,067 reference cells in this historical collection. It remains model-assisted review without a documented independent second reader. No reviewer credential or original review timestamp has been invented.')
    paragraph('The new training excludes all dense fields and their known source groups. Source grouping uses canonical SIPaKMeD field names and conservative APCData specimen-code proxies. A total of 340 noncanonical mirror fields and 5 APCData fields with unresolved source identity were quarantined. These identifiers support field/specimen grouping, not verified patient-level independence. Exact image hashes and group membership were checked across splits.')
    table(['Source','Training','Validation','Sparse test'],[['SIPaKMeD',648,139,139],['APCData',317,50,45]])
    paragraph('All unresolved-source dense fields were confined to development. Seeded canonical fields filled development to 20, leaving 20 canonical fields for evaluation. This deliberate source-composition difference must be considered when interpreting development-to-evaluation changes. The historical images had already been examined before this correction, so the reserved evaluation is not a previously unseen prospective cohort.')
    heading('Training and analysis protocol')
    paragraph('YOLOv8n was trained for each source variant at matched seeds 17, 43 and 101, with the same 150-epoch ceiling, patience 30, image size 640, automatic mixed precision and two loader workers. The installed Ultralytics fitness criterion is mAP50–95. Single-source and combined-source runs use different source compositions; their raw validation mAP values are not a direct superiority test.')
    paragraph('Each model threshold maximizes abnormal-cell F1 on the 20 development fields only. It is then frozen for dense evaluation and external testing. Matching is confidence-ordered greedy IoU of at least 0.5. Abnormal matching is class-specific; localization metrics separately ignore class. Percentile bootstrap intervals resample source-field groups or HMCHH patient-prefix proxies. They are conditional on fitted models and selected thresholds and do not establish patient-level clinical accuracy.')
    heading('Dense evaluation results')
    rows=[]
    for seed in (17,43,101):
        for variant in ('sipakmed_only','combined'):
            result=models[f'v2_{variant}_seed{seed}'];m=result['dense_evaluation']['metrics']
            rows.append(['Single' if variant=='sipakmed_only' else 'Combined',seed,f'{result["threshold_from_development"]:.2f}',f'{m["abnormal_precision"]:.3f}',f'{m["abnormal_recall"]:.3f}',f'{m["abnormal_f1"]:.3f}'])
    table(['Model','Seed','Threshold','Abn P','Abn R','Abn F1'],rows)
    for seed in (17,43,101):
        result=paired['paired_by_seed'][str(seed)]['dense'];interval=result['paired_cluster_bootstrap_95ci']['abnormal_recall']
        paragraph(f'Seed {seed}: combined minus single-source abnormal recall was {result["delta"]["abnormal_recall"]:+.3f}; conditional paired 95% bootstrap interval {interval[0]:+.3f} to {interval[1]:+.3f}.')
    paragraph('Full counts, false detections per field, localization metrics and uncertainty are preserved in the machine-readable result files. A cell-level recall is not patient sensitivity, and a detector precision is not a patient-level predictive value. Three-seed means, standard deviations and ranges are reported separately; no universal noise floor is inferred.')
    heading('External validation')
    first=next(iter(models.values()))['hmchh']
    paragraph(f'The complete available HMCHH validation split contributed {first["counts"]["fields"]} fields and {first["counts"]["abnormal_reference_cells"]} abnormal reference cells. None of these images were used in the new training. HMCHH annotates abnormal cells only, so abnormal predictions are assessed against abnormal references. Discarding predictions labeled Normal is not proof that those predictions represent correctly identified normal cells.')
    rows=[]
    for variant in ('sipakmed_only','combined'):
        stats=paired['seed_variability'][variant]['hmchh']
        rows.append(['Single' if variant=='sipakmed_only' else 'Combined',f'{stats["abnormal_precision"]["mean"]:.3f}',f'{stats["abnormal_recall"]["mean"]:.3f}',f'{stats["abnormal_f1"]["mean"]:.3f}',f'{stats["abnormal_false_detections_per_field"]["mean"]:.2f}'])
    table(['Model','Mean abn P','Mean abn R','Mean abn F1','False per field'],rows)
    paragraph('These means summarize three trained models, not three independent patient cohorts. APCData sparse-test results are also available, but combined-source models have seen APCData training images; their APCData results therefore are not zero-shot transfer. Incomplete source labels limit the interpretation of sparse-test precision.')
    heading('Calibration and normalization')
    paragraph('Temperature is fitted by binary negative log likelihood on development predictions and evaluated on separate dense evaluation predictions. ECE, Brier score and NLL are reported. This assesses confidence among detected candidates at a fixed 0.01 inference floor; it does not calibrate missed-cell risk or patient diagnosis.')
    rows=[]
    for variant in ('sipakmed_only','combined'):
        cal=[models[f'v2_{variant}_seed{s}']['calibration'] for s in (17,43,101)]
        rows.append(['Single' if variant=='sipakmed_only' else 'Combined',f'{mean(c["raw"]["ece"] for c in cal):.4f}',f'{mean(c["scaled"]["ece"] for c in cal):.4f}',f'{mean(c["raw"]["brier"] for c in cal):.4f}',f'{mean(c["scaled"]["brier"] for c in cal):.4f}'])
    table(['Model','Raw ECE','Scaled ECE','Raw Brier','Scaled Brier'],rows)
    deltas=[v['paired_difference']['delta']['abnormal_recall'] for model in stain['models'].values() for v in model['dense_evaluation']]
    paragraph(f'Reinhard inference-time normalization was tested using three distinct sets of 50 training-only reference images for each trained model. Across the 18 model/reference combinations, the dense-evaluation abnormal-recall change ranged from {min(deltas):+.3f} to {max(deltas):+.3f}. This measures reference and model sensitivity for this transform. It does not demonstrate morphological stain invariance or rule out other normalization methods. Thresholds were not retuned after transformation.')
    heading('Overall assessment and remaining limits')
    paragraph('The corrected work is a reproducible retrospective detector experiment with explicit data and evaluation limitations. Its useful contribution is examining annotation completeness, source composition and measurement reliability under a modest compute budget. It does not establish a novel architecture, a performance ceiling caused by limited VRAM, prospective patient-level effectiveness or a validated screening threshold.')
    paragraph('Before a clinical or confirmatory claim, obtain independently adjudicated annotations on a new patient-grouped cohort, evaluate a frozen protocol, and measure slide-level aggregation, missed-abnormal-slide burden and workflow consequences. Those future studies cannot be replaced by more epochs on the current fields.')
    heading('Sources and reproducibility')
    paragraph('Primary sources: SIPaKMeD original dataset at https://www.cs.uoi.gr/~marina/sipakmed.html; APCData doi:10.17632/ytd568rh3p.1; HMCHH dataset paper doi:10.1038/s41597-025-04374-5. The Coskun et al. classification study doi:10.3390/bioengineering13030289 is contextual literature, not a directly comparable detection benchmark.')
    paragraph('The research_v2 folder contains immutable protocol files, source and environment snapshots, split manifests, model hashes, per-field counts, prediction caches, bootstrap comparisons and normalization analyses. Historical results were archived before correction. The original audit and its verification evidence remain available for traceability.')
    footer=section.footer.paragraphs[0]
    footer.text='Corrected retrospective research report  |  '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    destination=ROOT/'results/Corrected_Research_Report.docx'
    doc.save(destination)
    (ROOT/'results/Corrected_Research_Report.txt').write_text('\n\n'.join(text),encoding='utf-8')
    print(destination)


if __name__ == '__main__':
    main()
